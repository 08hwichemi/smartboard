# 결석계 출력 양식(부광고 엑셀 "결석계" 파일의 인쇄 시트 2개)을 엑셀 모양 그대로 HTML 표로 옮겨 absence/forms.json을 만든다.
# 칸 너비·줄 높이·병합·테두리·글꼴·바탕색을 그대로 가져오고, 수식 칸은 비워 두고 data-c="B2"로 표시 → 화면(index.html)이 값을 채움.
# 쓰는 법: pip install openpyxl && python3 tools/absence_forms_build.py 결석계.xlsm absence/forms.json
# (원본 엑셀엔 학생 이름·결석 기록이 들어 있으니 이 공개 저장소에 올리지 말 것. 양식이 바뀌면 엑셀 인쇄 시트만 고치고 다시 돌리면 됨.)
import openpyxl, json, sys, html, re
from openpyxl.utils import get_column_letter as L, range_boundaries

THEME = ['FFFFFF', '000000', 'E7E6E6', '44546A', '4472C4', 'ED7D31', 'A5A5A5', 'FFC000', '5B9BD5', '70AD47']
MDW = 7.5  # 맑은 고딕 11 기본 글꼴 한 글자 폭(px) — 실제 엑셀 출력과 비교해 맞출 값

def color(c):
    if c is None: return None
    if c.type == 'rgb' and isinstance(c.rgb, str):
        return c.rgb[-6:]
    if c.type == 'theme':
        base = THEME[c.theme] if c.theme < len(THEME) else '000000'
        t = c.tint or 0
        ch = [int(base[i:i+2], 16) for i in (0, 2, 4)]
        ch = [round(v * (1 + t)) if t < 0 else round(v + (255 - v) * t) for v in ch]
        return ''.join('%02X' % v for v in ch)
    return None

BW = {'thin': '1px solid', 'medium': '2px solid', 'thick': '3px solid', 'hair': '1px solid',
      'dotted': '1px dotted', 'dashed': '1px dashed', 'double': '3px double', 'mediumDashed': '2px dashed'}

def side(s):
    if not s or not s.style: return None
    return BW.get(s.style, '1px solid') + ' #' + (color(s.color) or '000000')

def build(ws, rng, name):
    c1, r1, c2, r2 = range_boundaries(rng)
    merges = {}
    covered = set()
    for m in ws.merged_cells.ranges:
        if m.min_row > r2 or m.min_col > c2: continue
        merges[(m.min_row, m.min_col)] = m
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if (r, c) != (m.min_row, m.min_col): covered.add((r, c))
    widths = []
    for c in range(c1, c2 + 1):
        d = ws.column_dimensions[L(c)]
        w = d.width if d.width else (ws.sheet_format.defaultColWidth or 8.43)
        widths.append(0 if d.hidden else round(w * MDW))
    out = [f'<table class="xf" data-form="{name}" style="width:{sum(widths)}px">', '<colgroup>']
    out += [f'<col style="width:{w}px">' for w in widths]
    out.append('</colgroup>')
    dyn = []
    for r in range(r1, r2 + 1):
        rd = ws.row_dimensions[r]
        h = rd.height if rd.height else (ws.sheet_format.defaultRowHeight or 15)
        if rd.hidden: h = 0
        out.append(f'<tr style="height:{h}pt">')
        for c in range(c1, c2 + 1):
            if (r, c) in covered: continue
            cell = ws.cell(r, c)
            m = merges.get((r, c))
            rs = cs = 1
            if m:
                rs, cs = m.max_row - m.min_row + 1, m.max_col - m.min_col + 1
            # 병합 칸은 오른쪽/아래 테두리를 가장자리 칸에서 가져옴
            br = ws.cell(r, c + cs - 1).border.right
            bb = ws.cell(r + rs - 1, c).border.bottom
            st = []
            for k, s in (('left', cell.border.left), ('top', cell.border.top), ('right', br), ('bottom', bb)):
                v = side(s)
                if v: st.append(f'border-{k}:{v}')
            f = cell.font
            if f.name and f.name != '맑은 고딕': st.append(f"font-family:'{f.name}',serif" if '바탕' in f.name else f"font-family:'{f.name}'")
            if f.sz and f.sz != 11: st.append(f'font-size:{f.sz:g}pt')
            if f.b: st.append('font-weight:bold')
            fc = color(f.color)
            if fc and fc != '000000': st.append(f'color:#{fc}')
            if cell.fill and cell.fill.fill_type == 'solid':
                bg = color(cell.fill.fgColor)
                if bg and bg != 'FFFFFF': st.append(f'background:#{bg}')
            a = cell.alignment
            if a.horizontal in ('center', 'centerContinuous'): st.append('text-align:center')
            elif a.horizontal == 'right': st.append('text-align:right')
            elif a.horizontal == 'distributed': st.append('text-align:justify;text-align-last:justify')
            if a.vertical == 'top': st.append('vertical-align:top')
            elif a.vertical == 'bottom' or a.vertical is None: st.append('vertical-align:bottom')
            if a.wrap_text: st.append('white-space:pre-wrap')
            if a.indent: st.append(f'padding-left:{a.indent * 9}px')
            attrs = ' class="g"' if a.horizontal in (None, 'general') else ''
            if rs > 1: attrs += f' rowspan="{rs}"'
            if cs > 1: attrs += f' colspan="{cs}"'
            v = cell.value
            text = ''
            if isinstance(v, str) and v.startswith('='):
                attrs += f' data-c="{cell.coordinate}"'
                dyn.append(cell.coordinate)
            elif v is not None:
                text = html.escape(str(v))
            out.append(f'<td{attrs} style="{";".join(st)}">{text}</td>')
        out.append('</tr>')
    out.append('</table>')
    return '\n'.join(out), dyn

def close_frame(html):
    """바깥 테두리 아래 선 닫기 — 교외체험학습 시트는 양옆 선이 34행까지 내려오는데 아래 선이 엑셀 셀에
    안 남아 있어서 인쇄하면 큰 상자 아래가 열려 보였다(2026-09-30 사용자 제보). 첫 칸에 왼쪽 선이 있는
    마지막 줄의 칸들에 아래 선을 넣는다(이미 있으면 그대로)."""
    rows = re.findall(r'<tr[^>]*>.*?</tr>', html, re.S)
    framed = [r for r in rows if re.search(r'<td[^>]*style="[^"]*border-left', r.split('</td>')[0])]
    if not framed:
        return html
    last = framed[-1]
    if 'border-bottom' in last.split('</td>')[0]:
        return html
    fixed = re.sub(r'(<td[^>]*style=")(?![^"]*border-bottom)', r'\1border-bottom:1px solid #000000;', last)
    return html.replace(last, fixed)

wb = openpyxl.load_workbook(sys.argv[1])
res = {}
for sheet, rng, key in (('출력미리보기', 'A1:M31', 'confirm'), ('교외체험학습', 'A1:O35', 'trip')):
    ws = wb[sheet]
    t, dyn = build(ws, rng, key)
    t = close_frame(t)
    pm = ws.page_margins
    res[key] = {'html': t, 'dyn': dyn, 'scale': ws.page_setup.scale,
                'margin_in': [pm.top, pm.right, pm.bottom, pm.left], 'hcenter': ws.print_options.horizontalCentered}
json.dump(res, open(sys.argv[2], 'w'), ensure_ascii=False)
for k, v in res.items(): print(k, v['scale'], v['margin_in'], v['hcenter'], v['dyn'])
